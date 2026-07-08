"""Excel report — openpyxl workbook: Summary / Validations / RuleResults /
Mismatches, with status coloring, autofilter, and frozen headers."""

from __future__ import annotations

import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dvf.reporting.base import ReportGenerator, status_color
from dvf.reporting.data import ReportData
from dvf.utils.io import join_path, write_bytes

_HEADER_FILL = PatternFill("solid", fgColor="0F172A")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)


def _fill(status: str | None) -> PatternFill:
    return PatternFill("solid", fgColor=status_color(status).lstrip("#").upper())


def _status_font() -> Font:
    return Font(color="FFFFFF", bold=True, size=10)


class ExcelReportGenerator(ReportGenerator):
    format_name = "excel"

    def generate(self, data: ReportData, output_dir: str) -> str:
        wb = Workbook()
        self._summary_sheet(wb.active, data)
        self._table_sheet(
            wb.create_sheet("Validations"),
            ["validation_id", "status", "execution_group", "validation_profile",
             "source_label", "target_label", "attempt", "rules_passed",
             "rules_total", "source_row_count", "target_row_count",
             "rows_mismatched", "rows_missing", "rows_extra",
             "comparison_tier_used", "duration_sec", "error_message"],
            data.validations, status_col="status")
        self._table_sheet(
            wb.create_sheet("RuleResults"),
            ["validation_id", "rule_type", "column_name", "status", "source_value",
             "target_value", "difference", "difference_pct", "tolerance_applied",
             "duration_sec"],
            data.results, status_col="status")
        self._table_sheet(
            wb.create_sheet("Mismatches"),
            ["validation_id", "mismatch_type", "key_values", "column_name",
             "source_value", "target_value", "bucket_path"],
            data.mismatches, status_col=None)

        buffer = io.BytesIO()
        wb.save(buffer)
        path = join_path(output_dir, f"dvf_report_{data.run['run_id']}.xlsx")
        return write_bytes(path, buffer.getvalue())

    # ------------------------------------------------------------------ #
    def _summary_sheet(self, ws, data: ReportData) -> None:
        ws.title = "Summary"
        run = data.run
        rows = [
            ("Run ID", run.get("run_id")),
            ("Run name", run.get("run_name")),
            ("Status", run.get("status")),
            ("Trigger", run.get("trigger_type")),
            ("Triggered by", run.get("triggered_by")),
            ("Parent run", run.get("parent_run_id")),
            ("Started", str(run.get("start_ts") or "")),
            ("Finished", str(run.get("end_ts") or "")),
            ("Duration (s)", run.get("duration_sec")),
            ("Total validations", run.get("total_validations")),
            ("Passed", run.get("passed")),
            ("Failed", run.get("failed")),
            ("Warnings", run.get("warnings")),
            ("Errors", run.get("errors")),
            ("Skipped", run.get("skipped")),
            ("Mismatch rows captured", data.mismatch_total),
            ("Framework version", run.get("framework_version")),
            ("Report generated", data.generated_ts),
        ]
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 48
        for i, (label, value) in enumerate(rows, start=1):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True, size=10)
            cell = ws.cell(row=i, column=2, value=value)
            if label == "Status" and value:
                cell.fill = _fill(str(value))
                cell.font = _status_font()

    def _table_sheet(self, ws, columns: list[str], rows: list[dict],
                     status_col: str | None) -> None:
        for c, name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=c, value=name)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = \
                max(14, min(40, len(name) + 6))
        for r, row in enumerate(rows, start=2):
            for c, name in enumerate(columns, start=1):
                value = row.get(name)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, default=str)
                cell = ws.cell(row=r, column=c, value=value)
                if status_col and name == status_col and value:
                    cell.fill = _fill(str(value))
                    cell.font = _status_font()
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
